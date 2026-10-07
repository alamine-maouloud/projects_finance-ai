"""Tests de RiskLens. Lancer : python -m pytest

Les tests du facteur de spread utilisent une série de spread de test, générée ici et
clairement synthétique : elle sert à vérifier les calculs, jamais à produire des résultats.
"""
from __future__ import annotations

import csv
import json
import math
import shutil
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from risklens import extraction
from risklens.__main__ import analyser
from risklens.controles import anomalies_inventaire, concentration_lignes, dispersion_emetteurs
from risklens.donnees import (DATE_REFERENCE, DOSSIER_DONNEES, SOUS_TOTAL_TAUX_FIXE_EUR, charger_courbes,
                              empreinte, verifier_empreintes)
from risklens.excel import generer_classeur
from risklens.portefeuille import LIQUIDITES, Portefeuille
from risklens.risque import NOM_SPREAD, backtester, kupiec, mesurer, queue
from risklens.stress import stresser
from risklens.tableau_de_bord import generer_tableau

RACINE = Path(__file__).resolve().parent.parent
BUND = "DE0001102473"


@pytest.fixture(scope="module")
def reel():
    return analyser(avec_spread=False)


@pytest.fixture(scope="module")
def dossier_avec_spread(tmp_path_factory):
    """Copie des données réelles plus une série de spread de test (synthétique, pour les calculs seulement)."""
    dossier = tmp_path_factory.mktemp("donnees")
    for nom in ("inventory-fixed-eur.csv", "ecb-curves.csv", "manifeste.json"):
        shutil.copy(DOSSIER_DONNEES / nom, dossier / nom)
    dates = [d for d, _ in charger_courbes()]
    graine, lignes = 20250630, []
    for i, d in enumerate(dates):
        if i % 97 == 13:                       # quelques dates absentes pour tester la jointure stricte
            continue
        graine = (1664525 * graine + 1013904223) % 2**32
        bruit = graine / 2**32 * 2 - 1
        lignes.append((d.isoformat(), 0.0100 + 0.0040 * math.sin(i / 45) + 0.0006 * bruit))
    with (dossier / "spread-credit.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["date", "rendement_entreprises", "rendement_bund", "spread"])
        w.writerows((d, 0, 0, s) for d, s in lignes)
    manifeste = json.loads((dossier / "manifeste.json").read_text(encoding="utf-8"))
    manifeste["derives"]["spread-credit.csv"] = empreinte(dossier / "spread-credit.csv")
    (dossier / "manifeste.json").write_text(json.dumps(manifeste), encoding="utf-8")
    return dossier


@pytest.fixture(scope="module")
def avec_spread(dossier_avec_spread):
    return analyser(dossier=dossier_avec_spread)


# ---------------------------------------------------------------- non-régression et données

def test_reproduit_la_version_javascript_sur_l_ancienne_selection(reel):
    """Même sélection que la première version (19 lignes, filtre « 366 ») : mêmes résultats."""
    eligibles = [l for l in reel["inventaire"] if l.date_prix == DATE_REFERENCE and l.prix_clean > 0
                 and DATE_REFERENCE < l.echeance <= date(2035, 6, 30) and l.code_parenthese == "366"]
    selection = sorted(eligibles, key=lambda l: (-l.valeur, l.isin))[:19]
    pf = Portefeuille.depuis_lignes(selection, reel["courbes"][-1][1])
    m = mesurer(pf, reel["scenarios"], 0.99)
    s = stresser(pf, 100, 150)
    assert m.var == pytest.approx(4_753_665.44037919, abs=0.01)
    assert m.es == pytest.approx(6_005_584.06810319, abs=0.01)
    assert s.perte_lineaire == pytest.approx(76_137_658.7483302, abs=0.01)
    assert s.perte_revalorisee == pytest.approx(70_810_263, abs=1)


def test_resultats_de_reference_sur_l_univers_complet(reel):
    m = reel["mesure"]
    assert len(reel["retenues"]) == 167
    assert m.var == pytest.approx(17_023_014.53, abs=0.01)
    assert m.es == pytest.approx(21_571_463.67, abs=0.01)
    assert reel["spread_raccorde"] is False and m.contributions_facteurs[NOM_SPREAD] == 0


def test_univers_couvre_toute_la_section_valorisee(reel):
    assert len(reel["inventaire"]) == 169
    assert sum(l.valeur for l in reel["inventaire"]) == pytest.approx(SOUS_TOTAL_TAUX_FIXE_EUR, abs=0.02)
    assert {l.isin for l, _ in reel["exclues"]} == {"FR0012317451", "IT0005440646"}
    assert sum(l.valeur for l, _ in reel["exclues"]) == 0


def test_un_fichier_modifie_est_refuse(tmp_path):
    for nom in ("inventory-fixed-eur.csv", "ecb-curves.csv", "manifeste.json"):
        shutil.copy(DOSSIER_DONNEES / nom, tmp_path / nom)
    with (tmp_path / "ecb-curves.csv").open("a") as f:
        f.write("\n")
    with pytest.raises(ValueError, match="a changé"):
        verifier_empreintes(tmp_path)


def test_scenarios_successifs_et_anterieurs_a_la_date_de_reference(reel):
    sc = reel["scenarios"]
    assert len(sc) == 1000 and sc[-1].date == DATE_REFERENCE
    niveaux = dict(reel["courbes"])
    for s in sc:
        for k in range(3):
            assert s.variations[k] == pytest.approx(niveaux[s.date][k] - niveaux[s.date_precedente][k], abs=1e-15)
    assert all(a.date == b.date_precedente for a, b in zip(sc, sc[1:]))


def test_calibration_retrouve_le_prix_et_les_sensibilites(reel):
    for p in reel["portefeuille"].positions:
        flux = p.calibration.flux
        assert sum(f.valeur_actuelle for f in flux) == pytest.approx(p.ligne.prix_dirty, abs=1e-8)
        for k in range(3):
            eps = 1e-6
            def prix(choc):
                return sum(f.valeur_actuelle * math.exp(-choc * f.poids[k] * f.t) for f in flux)
            derivee = (prix(-eps) - prix(eps)) / (2 * eps * p.ligne.prix_dirty)
            assert derivee == pytest.approx(p.sensibilites[k], abs=1e-6)


# ---------------------------------------------------------------- mesures de risque

def test_queue_sur_un_exemple_calcule_a_la_main():
    q = queue([0, 1, 2, 3], 0.625)
    assert q.var == 2 and q.es == pytest.approx(8 / 3) and sum(q.poids) == pytest.approx(1)
    with pytest.raises(ValueError):
        queue([], 0.99)
    with pytest.raises(ValueError):
        queue([float("nan")], 0.99)


@pytest.mark.parametrize("confiance", [0.95, 0.99])
def test_contributions_se_rapprochent_de_l_es(reel, avec_spread, confiance):
    for res in (reel, avec_spread):
        pf = res["portefeuille"]
        for p in (pf, pf.vendre(BUND, 50e6), pf.vendre_prorata(100e6)):
            m = mesurer(p, res["scenarios"], confiance)
            assert sum(m.contributions_lignes.values()) == pytest.approx(m.es, abs=1e-4)
            assert sum(m.contributions_facteurs.values()) == pytest.approx(m.es, abs=1e-4)
            assert m.es >= m.var


def test_les_ventes_conservent_le_capital(reel):
    pf = reel["portefeuille"]
    apres = pf.vendre(BUND, 50e6)
    assert apres.valeur_totale == pytest.approx(pf.valeur_totale) and apres.montants[LIQUIDITES] == 50e6
    with pytest.raises(ValueError):
        pf.vendre(BUND, 1e12)


def test_vente_au_prorata_reduit_l_es_proportionnellement(reel):
    pf, sc = reel["portefeuille"], reel["scenarios"]
    base = mesurer(pf, sc).es
    assert mesurer(pf.vendre_prorata(100e6), sc).es == pytest.approx(base * (1 - 100e6 / pf.valeur_totale), rel=1e-12)


def test_backtest_n_utilise_aucune_donnee_future(reel):
    pf, sc = reel["portefeuille"], reel["scenarios"]
    reference = backtester(pf, sc)
    modifie = list(sc)
    modifie[-1] = replace(sc[-1], variations=(0.05, 0.05, 0.05))
    autre = backtester(pf, modifie)
    assert autre.var == reference.var and autre.pertes[-1] != reference.pertes[-1]


def test_kupiec_valeur_connue():
    lr, p = kupiec(3, 250, 0.99)
    assert lr == pytest.approx(0.0949401, abs=1e-6) and p == pytest.approx(0.7579883, abs=1e-6)


def test_stress_nul_convexe_et_souverain_sans_choc_credit(reel):
    pf = reel["portefeuille"]
    nul = stresser(pf, 0, 0)
    assert nul.perte_lineaire == 0 and nul.perte_revalorisee == pytest.approx(0, abs=1e-6)
    s = stresser(pf, 100, 150)
    assert 0 < s.perte_revalorisee < s.perte_lineaire
    assert s.perte_lineaire == pytest.approx(s.perte_taux + s.perte_credit)
    assert pf.position(BUND).duration_spread == 0


# ---------------------------------------------------------------- facteur de spread

def test_jointure_stricte_avec_le_spread(avec_spread, dossier_avec_spread):
    with (dossier_avec_spread / "spread-credit.csv").open() as f:
        spreads = {date.fromisoformat(r["date"]): float(r["spread"]) for r in csv.DictReader(f)}
    sc = avec_spread["scenarios"]
    assert len(sc) == 1000
    for s in sc:
        assert s.date in spreads and s.date_precedente in spreads
        assert s.variation_spread == pytest.approx(spreads[s.date] - spreads[s.date_precedente], abs=1e-15)


def test_beta_dts_et_methode_uniforme(dossier_avec_spread, avec_spread):
    pf, ref = avec_spread["portefeuille"], avec_spread["spread_reference"]
    for p in pf.positions:
        assert p.beta_spread == pytest.approx(max(p.calibration.spread, 0) / ref)
    assert pf.position(BUND).sensibilite_spread == 0
    uniforme = analyser(dossier=dossier_avec_spread, methode_spread="uniforme")["portefeuille"]
    assert all(p.beta_spread == 1 for p in uniforme.positions)


def test_le_spread_ajoute_du_risque(reel, avec_spread):
    assert avec_spread["mesure"].contributions_facteurs[NOM_SPREAD] != 0
    assert reel["mesure"].contributions_facteurs[NOM_SPREAD] == 0


# ---------------------------------------------------------------- contrôles

def test_controles_detectent_les_anomalies_reelles(reel):
    alertes = anomalies_inventaire(reel["inventaire"], reel["portefeuille"])
    par_ligne = {}
    for a in alertes:
        par_ligne.setdefault(a.isin, set()).add(a.controle)
    assert par_ligne["FR0012317451"] >= {"Titre échu toujours en inventaire", "Valorisation nulle", "Prix ancien"}
    assert "Valorisation nulle" in par_ligne["IT0005440646"]
    assert {a.isin for a in alertes if a.controle == "Spread élevé"} == {
        "XS2587104444", "XS2393687350", "XS2919072962", "XS2835773255"}
    assert concentration_lignes(reel["portefeuille"]) == []
    d = dispersion_emetteurs(reel["inventaire"])
    assert d.plus_de_5 == {} and d.depassements_10 == {}


# ---------------------------------------------------------------- extraction

def test_courbes_bce_reconstruites_a_l_identique():
    courbes = extraction.lire_courbes_bce([extraction.BRUT / "ecb-5y.csv", extraction.BRUT / "ecb-2y-10y.csv"])
    with (DOSSIER_DONNEES / "ecb-curves.csv").open() as f:
        existantes = list(csv.DictReader(f))
    assert len(courbes) == len(existantes)
    for c, e in zip(courbes, existantes):
        assert c[0] == e["date"] and c[1:] == (float(e["rate2"]), float(e["rate5"]), float(e["rate10"]))


def test_lecture_de_la_section_taux_fixe_sur_texte_reconstitue():
    """Le texte des pages est reconstitué au format de l'inventaire à partir des 169 lignes extraites."""
    with (DOSSIER_DONNEES / "inventory-fixed-eur.csv").open() as f:
        lignes = list(csv.DictReader(f))
    def texte(r):
        d = date.fromisoformat(r["maturity"])
        q = date.fromisoformat(r["quoteDate"])
        return (f"{r['id']} {r['name']} ({r['codeParenthese']}) {d:%d%m%y} {float(r['nominal']):,.2f} M EUR "
                f"{100:,.4f} % {q:%d/%m/%y} {float(r['cleanPrice']):,.4f} A {0:,.2f} {float(r['marketValue']):,.2f} "
                f"{float(r['accruedInterest']):,.2f} {0:,.2f} {0:,.2f}")
    pages = ["Titre\n" + extraction.DEBUT_SECTION + "\n" + "\n".join(texte(r) for r in lignes) + "\n" +
             extraction.FIN_SECTION + "\nXS0000000000 hors section"]
    lus = extraction.lire_section_taux_fixe(pages)
    assert len(lus) == 169
    for lu, r in zip(lus, lignes):
        assert lu["id"] == r["id"] and lu["maturity"] == r["maturity"] and lu["codeParenthese"] == r["codeParenthese"]
        assert lu["marketValue"] == pytest.approx(float(r["marketValue"]), abs=0.005)
        assert lu["coupon"] == pytest.approx(float(r["coupon"]))


@pytest.mark.parametrize("contenu", [
    '"";"BBSIS.D";"BBSIS.D_FLAGS"\n"";"Rendite";\n"Einheit";"% p.a.";\n2025-06-27;"3,10";\n2025-06-30;"3,12";\n2025-07-01;".";\n',
    "TIME_PERIOD,OBS_VALUE,OBS_STATUS\n2025-06-27,3.10,A\n2025-06-30,3.12,A\n",
])
def test_lecture_des_series_bundesbank(tmp_path, contenu):
    chemin = tmp_path / "serie.csv"
    chemin.write_text(contenu, encoding="utf-8")
    assert extraction.lire_serie_bundesbank(chemin) == {"2025-06-27": 3.10, "2025-06-30": 3.12}


def test_construction_du_spread_et_controles_de_plausibilite():
    entreprises = {"2025-06-27": 3.10, "2025-06-30": 3.12, "2025-07-01": 3.0}
    bund = {"2025-06-27": 2.30, "2025-06-30": 2.35}
    lignes = extraction.construire_spread(entreprises, bund, {"2025-06-27", "2025-06-30"})
    assert [l[0] for l in lignes] == ["2025-06-27", "2025-06-30"]
    assert lignes[-1][3] == pytest.approx(0.0077)
    with pytest.raises(ValueError):
        extraction.construire_spread({"2025-06-30": 45.0}, {"2025-06-30": 2.35}, {"2025-06-30"})


# ---------------------------------------------------------------- classeur, tableau de bord, présentation

CONFIGURATIONS = [(0.99, None, 20e6, 100, 150), (0.95, BUND, 50e6, 200, 0)]


def _configurations(res):
    return [(c, isin or res["isin_vente"], m, t, s) for c, isin, m, t, s in CONFIGURATIONS]


@pytest.mark.parametrize("cas", ["reel", "avec_spread"])
def test_classeur_concorde_avec_le_moteur(tmp_path, cas, request):
    from risklens import verification_excel as v
    if v.trouver_soffice() is None:
        pytest.skip("LibreOffice absent : vérification du classeur non exécutée")
    res = request.getfixturevalue(cas)
    chemin = tmp_path / "RiskLens-Comite.xlsx"
    generer_classeur(chemin, res)
    assert v.comparer(chemin, res["portefeuille"], res["scenarios"], _configurations(res)) == []


@pytest.mark.parametrize("cas", ["reel", "avec_spread"])
def test_tableau_de_bord_concorde_avec_le_moteur(tmp_path, cas, request):
    from risklens import verification_tableau as v
    if not v.playwright_disponible():
        pytest.skip("Playwright absent : vérification du tableau de bord non exécutée")
    res = request.getfixturevalue(cas)
    chemin = tmp_path / "RiskLens.html"
    generer_tableau(chemin, res)
    assert v.comparer(chemin, res["portefeuille"], res["scenarios"], _configurations(res)) == []


def test_aucun_tiret_long_dans_le_projet(tmp_path, reel):
    interdits = ("\u2014", "\u2013")
    fichiers = [f for f in RACINE.rglob("*") if f.is_file() and "sorties" not in f.parts and not any(p.startswith(".") for p in f.relative_to(RACINE).parts)
                and f.suffix in {".py", ".md", ".bas", ".json", ".toml"}]
    for f in fichiers:
        texte = f.read_text(encoding="utf-8", errors="ignore")
        assert not any(c in texte for c in interdits), f"Tiret long dans {f.relative_to(RACINE)}"
    generer_classeur(tmp_path / "classeur.xlsx", reel)
    generer_tableau(tmp_path / "tableau.html", reel)
    html = (tmp_path / "tableau.html").read_text(encoding="utf-8")
    assert not any(c in html for c in interdits), "Tiret long dans le tableau de bord"
    wb = load_workbook(tmp_path / "classeur.xlsx")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not any(x in c.value for x in interdits), f"Tiret long en {ws.title}!{c.coordinate}"
                assert not any(x in (c.number_format or "") for x in interdits)
