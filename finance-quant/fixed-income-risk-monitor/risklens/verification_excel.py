"""Vérifie que le classeur Excel recalculé donne les mêmes résultats que le moteur Python.

Le classeur est modifié (paramètres), recalculé par LibreOffice en mode sans interface,
puis relu. Utilisation : python -m risklens.verification_excel
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from .portefeuille import Portefeuille
from .risque import mesurer
from .stress import stresser

CHEMINS_SOFFICE = ["soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"]


def trouver_soffice() -> str | None:
    for c in CHEMINS_SOFFICE:
        if shutil.which(c) or Path(c).exists():
            return shutil.which(c) or c
    return None


def recalculer(source: Path, dossier: Path) -> Path:
    """Ouvre le classeur dans LibreOffice, qui calcule toutes les formules, et l'enregistre."""
    soffice = trouver_soffice()
    if soffice is None:
        raise FileNotFoundError("LibreOffice introuvable : installez-le pour vérifier le classeur.")
    with tempfile.TemporaryDirectory(prefix="lo_profil_") as profil:
        subprocess.run([soffice, f"-env:UserInstallation={Path(profil).as_uri()}", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", str(dossier), str(source)],
                       check=True, capture_output=True, timeout=600, env=os.environ.copy())
    return dossier / source.name


def valeurs_classeur(chemin: Path, confiance: float, isin: str, montant: float,
                     choc_taux: float, choc_spread: float) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        wb = load_workbook(chemin)
        noms = wb.defined_names
        def cible(nom):
            feuille, cellule = noms[nom].attr_text.rsplit("!", 1)
            return wb[feuille.strip("'").replace("''", "'")][cellule.replace("$", "")]
        for nom, valeur in (("RL_Confiance", confiance), ("RL_ISIN_Vente", isin), ("RL_Montant_Vente", montant),
                            ("RL_Choc_Taux", choc_taux), ("RL_Choc_Spread", choc_spread)):
            cible(nom).value = valeur
        entree = tmp / "entree" / chemin.name
        entree.parent.mkdir()
        wb.save(entree)
        sortie = recalculer(entree, tmp)
        calc = load_workbook(sortie, data_only=True)
        def lire(nom):
            feuille, cellule = calc.defined_names[nom].attr_text.rsplit("!", 1)
            return calc[feuille.strip("'").replace("''", "'")][cellule.replace("$", "")].value
        com = calc["Comité"]
        erreurs = [f"{ws.title}!{c.coordinate}" for ws in calc.worksheets for row in ws.iter_rows()
                   for c in row if isinstance(c.value, str) and c.value.startswith("#")]
        return dict(var_initial=lire("RL_VaR_Initial"), var_apres=lire("RL_VaR_Apres"),
                    es_initial=lire("RL_ES_Initial"), es_apres=lire("RL_ES_Apres"),
                    stress_initial=com["C13"].value, stress_apres=com["D13"].value,
                    reval_initial=com["C14"].value, reval_apres=com["D14"].value,
                    valide=lire("RL_Valide"), erreurs=erreurs)


def valeurs_python(portefeuille: Portefeuille, scenarios, confiance, isin, montant, choc_taux, choc_spread) -> dict:
    apres = portefeuille.vendre(isin, montant)
    m0, m1 = mesurer(portefeuille, scenarios, confiance), mesurer(apres, scenarios, confiance)
    s0, s1 = stresser(portefeuille, choc_taux, choc_spread), stresser(apres, choc_taux, choc_spread)
    return dict(var_initial=m0.var, var_apres=m1.var, es_initial=m0.es, es_apres=m1.es,
                stress_initial=s0.perte_lineaire, stress_apres=s1.perte_lineaire,
                reval_initial=s0.perte_revalorisee, reval_apres=s1.perte_revalorisee)


def comparer(chemin: Path, portefeuille, scenarios, configurations, tolerance=0.01) -> list[str]:
    """Renvoie la liste des écarts (vide si le classeur et le moteur concordent)."""
    ecarts = []
    for conf in configurations:
        excel = valeurs_classeur(chemin, *conf)
        python = valeurs_python(portefeuille, scenarios, *conf)
        if excel["erreurs"]:
            ecarts.append(f"{conf} : erreurs de formule {excel['erreurs'][:5]}")
        if excel["valide"] != "OK":
            ecarts.append(f"{conf} : paramètres jugés invalides par le classeur")
        for cle, v in python.items():
            if excel[cle] is None or abs(excel[cle] - v) > tolerance:
                ecarts.append(f"{conf} : {cle} Excel {excel[cle]} contre Python {v}")
    return ecarts


if __name__ == "__main__":
    from .__main__ import SORTIES, analyser
    r = analyser()
    pf, sc = r["portefeuille"], r["scenarios"]
    bund = "DE0001102473"
    configurations = [(0.99, r["isin_vente"], 20e6, 100, 150), (0.95, r["isin_vente"], 20e6, 100, 150),
                      (0.99, bund, 50e6, 200, 0), (0.99, bund, 0, -100, 250)]
    ecarts = comparer(SORTIES / "RiskLens-Comite.xlsx", pf, sc, configurations)
    print("Classeur et moteur concordent." if not ecarts else "\n".join(ecarts))
