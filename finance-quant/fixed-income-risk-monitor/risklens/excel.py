"""Génération du classeur Excel du comité des risques.

Le classeur recalcule tout par formules à partir des positions, des sensibilités
calibrées par le moteur Python et des variations de taux observées. Il s'adapte
au nombre de lignes de l'univers. Les plages nommées « RL_... » servent au module VBA.
"""
from __future__ import annotations

from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from .controles import SEUILS
from .donnees import (ACTIF_NET_FONDS, DATE_REFERENCE, SOUS_TOTAL_TAUX_FIXE_EUR, SOUVERAINS,
                      charger_manifeste, motifs_exclusion)
from .portefeuille import LIQUIDITES

POLICE = "Arial"
BLEU_FONCE = "19334A"
F_TITRE = Font(name=POLICE, size=14, bold=True, color=BLEU_FONCE)
F_SOUS_TITRE = Font(name=POLICE, size=9, italic=True, color="5B6B7A")
F_ENTETE = Font(name=POLICE, size=9, bold=True, color="FFFFFF")
F_NORMAL = Font(name=POLICE, size=9)
F_GRAS = Font(name=POLICE, size=9, bold=True)
F_SAISIE = Font(name=POLICE, size=9, color="0000FF")
F_NOTE = Font(name=POLICE, size=8, italic=True, color="5B6B7A")
FOND_ENTETE = PatternFill("solid", fgColor=BLEU_FONCE)
FOND_SAISIE = PatternFill("solid", fgColor="FFF2CD")
FOND_LEGER = PatternFill("solid", fgColor="EEF2F6")
FOND_ALERTE = PatternFill("solid", fgColor="F8D7DA")
FOND_EXAMEN = PatternFill("solid", fgColor="FFF3CD")
TRAIT = Border(bottom=Side(style="thin", color="C9D2DB"))

EUR = '#,##0;(#,##0);"-"'
EUR2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.00%;(0.00%);"-"'
DEC = '0.000;(0.000);"-"'
PB = '#,##0;(#,##0);"-"'
DATE = "dd/mm/yyyy"
TAUX = '0.000000%'

N_SCENARIOS = 1000
L0 = 7  # première ligne de données dans les onglets tabulaires


def _q(feuille: str) -> str:
    return "'" + feuille.replace("'", "''") + "'"


def _ref(feuille: str, cellule: str) -> str:
    return f"{_q(feuille)}!{cellule}"


def _nommer(wb, nom: str, feuille: str, plage: str) -> None:
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"{_q(feuille)}!{plage}")


def _entete(ws, ligne: int, colonne: int, libelles: list[str], largeurs: list[float] | None = None):
    for i, texte in enumerate(libelles):
        c = ws.cell(ligne, colonne + i, texte)
        c.font, c.fill = F_ENTETE, FOND_ENTETE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[ligne].height = 30
    if largeurs:
        for i, w in enumerate(largeurs):
            ws.column_dimensions[get_column_letter(colonne + i)].width = w


def _titre(ws, titre: str, sous_titre: str = "") -> None:
    ws["B2"] = titre
    ws["B2"].font = F_TITRE
    if sous_titre:
        ws["B3"] = sous_titre
        ws["B3"].font = F_SOUS_TITRE
    ws.column_dimensions["A"].width = 2


def _cellule(ws, ref: str, valeur, fmt: str | None = None, police=F_NORMAL, fond=None):
    c = ws[ref]
    c.value = valeur
    c.font = police
    if fmt:
        c.number_format = fmt
    if fond:
        c.fill = fond
    return c


def generer_classeur(chemin, res: dict) -> None:
    """Écrit le classeur à partir du résultat de risklens.__main__.analyser."""
    inventaire, portefeuille, scenarios = res["inventaire"], res["portefeuille"], res["scenarios"]
    spread_ok, methode_spread, indisponibles = res["spread_raccorde"], res["methode_spread"], res["indisponibles"]
    empreintes = res["empreintes"]
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel recalcule tout à l'ouverture
    positions = list(portefeuille.positions)
    n_pos = len(positions)
    p_fin = L0 + n_pos - 1          # dernière ligne de titre
    p_liq = p_fin + 1               # ligne des liquidités
    p_tot = p_liq + 1               # ligne des totaux
    h_fin = L0 + N_SCENARIOS - 1
    flux = [(p.isin, f) for p in positions for f in p.calibration.flux]
    f_fin = L0 + len(flux) - 1
    i_fin = L0 + len(inventaire) - 1

    COM, PAR, POS, CAL, HIS, FLX, BKT, CTL, INV, EMT, SRC = (
        "Comité", "Paramètres", "Positions", "Calculs", "Historique", "Flux",
        "Backtest", "Contrôles", "Inventaire", "Émetteurs", "Sources")
    ws_com = wb.active
    ws_com.title = COM
    feuilles = {COM: ws_com}
    for nom in (PAR, POS, CAL, HIS, FLX, BKT, CTL, INV, EMT, SRC):
        feuilles[nom] = wb.create_sheet(nom)
    for ws in feuilles.values():
        ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------ Paramètres
    ws = feuilles[PAR]
    _titre(ws, "Paramètres", "Cellules jaunes : saisie. Les seuils sont des seuils de revue du prototype, "
                             "pas des limites du fonds.")
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 70
    _entete(ws, 5, 2, ["Paramètre", "Valeur", "Commentaire"])
    parametres = [
        ("RL_Confiance", "Niveau de confiance", 0.99, PCT, "95 % ou 99 %"),
        ("RL_Choc_Taux", "Choc parallèle de taux (pb)", 100, PB, "De -300 à +300 pb"),
        ("RL_Choc_Spread", "Choc uniforme de spreads (pb)", 150, PB, "De -200 à +500 pb ; non appliqué au souverain AAA"),
        ("RL_ISIN_Vente", "Ligne à vendre (ISIN)", res["isin_vente"], None, "Liste déroulante ; vente réinvestie en liquidités, sans frais"),
        ("RL_Montant_Vente", "Montant de la vente (EUR)", res["montant_vente"], EUR, "Entre 0 et la valeur détenue sur la ligne"),
        ("RL_Spread_Revue", "Seuil de spread de revue (pb)", SEUILS["spread_revue_pb"], PB, "Au-delà : suivi crédit"),
        ("RL_Age_Prix", "Ancienneté maximale d'un prix (jours)", SEUILS["age_prix_jours"], "0", "Par rapport à la date de référence"),
        ("RL_Seuil_Ligne", "Seuil de revue par ligne", SEUILS["ligne_max"], PCT, "En % du portefeuille modélisé"),
        ("RL_Seuil_Emetteur_5", "Seuil émetteur, règle 5/10/40", SEUILS["emetteur_seuil"], PCT, "En % de l'actif net publié"),
        ("RL_Seuil_Emetteur_10", "Plafond émetteur, règle 5/10/40", SEUILS["emetteur_max"], PCT, "En % de l'actif net publié"),
        ("RL_Seuil_Cumul_40", "Plafond du cumul des émetteurs > 5 %", SEUILS["emetteurs_cumul_max"], PCT, "En % de l'actif net publié"),
    ]
    for i, (nom, libelle, valeur, fmt, commentaire) in enumerate(parametres):
        r = 6 + i
        _cellule(ws, f"B{r}", libelle)
        _cellule(ws, f"C{r}", valeur, fmt, F_SAISIE, FOND_SAISIE).alignment = Alignment(horizontal="right")
        _cellule(ws, f"D{r}", commentaire, police=F_NOTE)
        _nommer(wb, nom, PAR, f"$C${r}")
    r = 6 + len(parametres) + 1
    _cellule(ws, f"B{r}", "Date de référence", police=F_GRAS)
    _cellule(ws, f"C{r}", DATE_REFERENCE, DATE)
    _nommer(wb, "RL_Date", PAR, f"$C${r}")
    _cellule(ws, f"B{r + 1}", "Actif net publié du fonds (EUR)", police=F_GRAS)
    _cellule(ws, f"C{r + 1}", ACTIF_NET_FONDS, EUR)
    _nommer(wb, "RL_Actif_Net", PAR, f"$C${r + 1}")
    _cellule(ws, f"B{r + 2}", "Validité des paramètres", police=F_GRAS)
    _cellule(ws, f"C{r + 2}",
             "=IF(AND(OR(RL_Confiance=0.95,RL_Confiance=0.99),RL_Choc_Taux>=-300,RL_Choc_Taux<=300,"
             "RL_Choc_Spread>=-200,RL_Choc_Spread<=500,ISNUMBER(MATCH(RL_ISIN_Vente,RL_Liste_ISIN,0)),"
             "RL_Montant_Vente>=0,RL_Montant_Vente<=IFERROR(INDEX(" + _ref(POS, f"$J${L0}:$J${p_fin}") +
             ",MATCH(RL_ISIN_Vente,RL_Liste_ISIN,0)),0),RL_Spread_Revue>0,RL_Age_Prix>=0,"
             "RL_Seuil_Ligne>0,RL_Seuil_Emetteur_5>0,RL_Seuil_Emetteur_10>0,RL_Seuil_Cumul_40>0),\"OK\",\"À CORRIGER\")",
             police=F_GRAS)
    _nommer(wb, "RL_Valide", PAR, f"$C${r + 2}")
    ws.conditional_formatting.add(f"C{r + 2}", CellIsRule(operator="notEqual", formula=['"OK"'], fill=FOND_ALERTE))
    dv_conf = DataValidation(type="list", formula1='"0.95,0.99"', allow_blank=False)
    dv_isin = DataValidation(type="list", formula1="RL_Liste_ISIN", allow_blank=False)
    ws.add_data_validation(dv_conf)
    ws.add_data_validation(dv_isin)
    dv_conf.add("C6")
    dv_isin.add("C9")
    _cellule(ws, f"B{r + 4}", "Les chocs de stress sont hypothétiques et sans probabilité attribuée. "
                              "Le module VBA facultatif ajoute une vérification indépendante et les exports.",
             police=F_NOTE)

    # ------------------------------------------------------------------ Positions
    ws = feuilles[POS]
    _titre(ws, "Positions modélisées",
           f"{n_pos} lignes de la section taux fixe EUR. Valeurs publiées ; spreads et sensibilités "
           "estimés par le modèle de flux (valeurs importées du moteur Python).")
    colonnes = ["ISIN", "Libellé", "Émetteur", "Échéance", "Coupon (%)", "Nominal (EUR)", "Prix clean",
                "Coupon couru (EUR)", "Valeur publiée (EUR)", "Prix dirty", "Spread calibré (pb)",
                "Sensibilité 2 ans", "Sensibilité 5 ans", "Sensibilité 10 ans", "Duration",
                "Duration de spread", "Souverain AAA", "Montant après vente (EUR)", "Poids initial",
                "Poids après vente", "Contribution ES initiale (EUR)", "Contribution ES après (EUR)",
                "ES par M EUR détenu", "Stress linéaire initial (EUR)", "Stress linéaire après (EUR)",
                "Stress revalorisé initial (EUR)", "Stress revalorisé après (EUR)", "Au-dessus du seuil de ligne",
                "Écart flux / prix dirty", "Bêta de spread", "Sensibilité au facteur spread"]
    _entete(ws, L0 - 1, 2, colonnes, [14, 24, 14, 11, 9, 14, 9, 13, 15, 9, 10, 10, 10, 10, 9, 10, 9,
                                       15, 9, 9, 14, 14, 12, 14, 14, 14, 14, 11, 11, 9, 11])
    nav = "RL_NAV"
    choc_taux_queue_i = [_ref(CAL, f"$I${k}") for k in (13, 14, 15, 22)]
    choc_taux_queue_a = [_ref(CAL, f"$J${k}") for k in (13, 14, 15, 22)]
    flx_isin = _ref(FLX, f"$B${L0}:$B${f_fin}")
    flx_t = _ref(FLX, f"$D${L0}:$D${f_fin}")
    flx_va = _ref(FLX, f"$F${L0}:$F${f_fin}")
    for i, p in enumerate(positions):
        r = L0 + i
        l = p.ligne
        valeurs = [l.isin, l.libelle, l.emetteur, l.echeance, l.coupon, l.nominal, l.prix_clean,
                   l.coupon_couru, l.valeur]
        fmts = [None, None, None, DATE, "0.000", EUR, "0.0000", EUR2, EUR2]
        for j, (v, fmt) in enumerate(zip(valeurs, fmts)):
            _cellule(ws, f"{get_column_letter(2 + j)}{r}", v, fmt)
        _cellule(ws, f"K{r}", f"=J{r}/G{r}*100", "0.0000")
        _cellule(ws, f"L{r}", p.calibration.spread * 1e4, "0.0", F_SAISIE)
        for k, col in enumerate("MNO"):
            _cellule(ws, f"{col}{r}", p.sensibilites[k], "0.0000", F_SAISIE)
        _cellule(ws, f"P{r}", f"=SUM(M{r}:O{r})", "0.0000")
        _cellule(ws, f"R{r}", "OUI" if p.isin in SOUVERAINS else "NON", None, F_SAISIE)
        _cellule(ws, f"Q{r}", f'=IF(R{r}="OUI",0,P{r})', "0.0000")
        _cellule(ws, f"S{r}", f"=IF(B{r}=RL_ISIN_Vente,J{r}-RL_Montant_Vente,J{r})", EUR)
        _cellule(ws, f"T{r}", f"=J{r}/{nav}", "0.0000%")
        _cellule(ws, f"U{r}", f"=S{r}/{nav}", "0.0000%")
        _cellule(ws, f"V{r}", f"={nav}*T{r}*(M{r}*{choc_taux_queue_i[0]}+N{r}*{choc_taux_queue_i[1]}+O{r}*{choc_taux_queue_i[2]}+AF{r}*{choc_taux_queue_i[3]})", EUR)
        _cellule(ws, f"W{r}", f"={nav}*U{r}*(M{r}*{choc_taux_queue_a[0]}+N{r}*{choc_taux_queue_a[1]}+O{r}*{choc_taux_queue_a[2]}+AF{r}*{choc_taux_queue_a[3]})", EUR)
        _cellule(ws, f"X{r}", f"=IF(S{r}>0,W{r}/S{r}*1000000,0)", EUR)
        _cellule(ws, f"Y{r}", f"={nav}*T{r}*(P{r}*RL_Choc_Taux+Q{r}*RL_Choc_Spread)/10000", EUR)
        _cellule(ws, f"Z{r}", f"={nav}*U{r}*(P{r}*RL_Choc_Taux+Q{r}*RL_Choc_Spread)/10000", EUR)
        choc = f"(RL_Choc_Taux+IF(Q{r}>0,RL_Choc_Spread,0))/10000"
        reval = f"(1-SUMPRODUCT(({flx_isin}=B{r})*{flx_va}*EXP(-{choc}*{flx_t}))/K{r})"
        _cellule(ws, f"AA{r}", f"={nav}*T{r}*{reval}", EUR)
        _cellule(ws, f"AB{r}", f"={nav}*U{r}*{reval}", EUR)
        _cellule(ws, f"AC{r}", f'=IF(U{r}>RL_Seuil_Ligne,"OUI","")')
        _cellule(ws, f"AD{r}", f"=SUMIF({flx_isin},B{r},{flx_va})-K{r}", "0.000000000")
        _cellule(ws, f"AE{r}", p.beta_spread, "0.000", F_SAISIE)
        _cellule(ws, f"AF{r}", f"=Q{r}*AE{r}", "0.0000")
    r = p_liq
    _cellule(ws, f"B{r}", LIQUIDITES, police=F_GRAS)
    _cellule(ws, f"C{r}", "Liquidités issues des ventes simulées")
    _cellule(ws, f"J{r}", 0, EUR2)
    _cellule(ws, f"S{r}", "=RL_Montant_Vente", EUR)
    _cellule(ws, f"T{r}", f"=J{r}/{nav}", "0.0000%")
    _cellule(ws, f"U{r}", f"=S{r}/{nav}", "0.0000%")
    for col in ("M", "N", "O", "P", "Q", "AE", "AF"):
        _cellule(ws, f"{col}{r}", 0, "0.0000")
    for col in ("V", "W", "Y", "Z", "AA", "AB"):
        _cellule(ws, f"{col}{r}", 0, EUR)
    r = p_tot
    _cellule(ws, f"C{r}", "Total", police=F_GRAS)
    for col, fmt in (("J", EUR2), ("S", EUR), ("T", "0.0000%"), ("U", "0.0000%"), ("V", EUR), ("W", EUR),
                     ("Y", EUR), ("Z", EUR), ("AA", EUR), ("AB", EUR)):
        _cellule(ws, f"{col}{r}", f"=SUM({col}{L0}:{col}{p_liq})", fmt, F_GRAS)
    for col in range(2, 30):
        ws.cell(r, col).border = Border(top=Side(style="thin", color=BLEU_FONCE))
    _nommer(wb, "RL_NAV", POS, f"$J${p_tot}")
    _nommer(wb, "RL_Liste_ISIN", POS, f"$B${L0}:$B${p_fin}")
    _nommer(wb, "RL_Poids_Apres", POS, f"$U${L0}:$U${p_liq}")
    _nommer(wb, "RL_Sensibilites", POS, f"$M${L0}:$O${p_liq}")
    _nommer(wb, "RL_Sensibilite_Spread", POS, f"$AF${L0}:$AF${p_liq}")
    _nommer(wb, "RL_Tableau_Positions", POS, f"$B${L0 - 1}:$AC${p_liq}")
    ws.freeze_panes = f"D{L0}"
    ws.conditional_formatting.add(f"AC{L0}:AC{p_fin}", CellIsRule(operator="equal", formula=['"OUI"'], fill=FOND_EXAMEN))
    ws.conditional_formatting.add(f"B{L0}:C{p_fin}", FormulaRule(formula=[f"$B{L0}=RL_ISIN_Vente"], fill=FOND_LEGER, font=F_GRAS))

    # ------------------------------------------------------------------ Historique
    ws = feuilles[HIS]
    _titre(ws, "Historique des variations de taux et de spread",
           "Courbe AAA de la BCE, zéro-coupon continu ; spread crédit Bundesbank si raccordé. Variation entre deux "
           "dates d'observation communes successives, sans remplissage. Décimales.")
    _entete(ws, L0 - 1, 2, ["Date", "Date précédente", "Variation 2 ans", "Variation 5 ans", "Variation 10 ans",
                            "Variation du spread" if spread_ok else "Spread (non raccordé)"], [12, 14, 15, 15, 15, 17])
    for i, sc in enumerate(scenarios):
        r = L0 + i
        _cellule(ws, f"B{r}", sc.date, DATE)
        _cellule(ws, f"C{r}", sc.date_precedente, DATE)
        for k, col in enumerate("DEF"):
            _cellule(ws, f"{col}{r}", sc.variations[k], TAUX, F_SAISIE)
        _cellule(ws, f"G{r}", sc.variation_spread, TAUX, F_SAISIE)
    _nommer(wb, "RL_Variations", HIS, f"$D${L0}:$G${h_fin}")
    ws.freeze_panes = f"B{L0}"

    # ------------------------------------------------------------------ Calculs
    ws = feuilles[CAL]
    _titre(ws, "Distribution des pertes historiques",
           "Perte du jour = valeur x (sensibilités aux taux x variations de taux + sensibilité spread x variation "
           "du spread). Pertes positives. Poids de queue : part de chaque scénario dans l'ES.")
    _entete(ws, L0 - 1, 2, ["Date", "Perte initiale (EUR)", "Perte après vente (EUR)", "Poids de queue initial",
                            "Poids de queue après"], [12, 16, 16, 14, 14])
    _entete(ws, L0 - 1, 8, ["Grandeur", "Initial", "Après vente"], [34, 16, 16])
    resume = {
        7: ("Nombre de scénarios", f"=COUNT({_ref(HIS, f'$D${L0}:$D${h_fin}')})",
            f"=COUNT({_ref(HIS, f'$D${L0}:$D${h_fin}')})", "0"),
        8: ("Masse de queue (journées)", "=ROUND(I7*(1-RL_Confiance),8)", "=ROUND(J7*(1-RL_Confiance),8)", "0.00"),
        9: ("Seuil de queue (EUR)", f"=LARGE(C{L0}:C{h_fin},ROUNDUP(I8,0))", f"=LARGE(D{L0}:D{h_fin},ROUNDUP(J8,0))", EUR),
        10: ("Sensibilité 2 ans du portefeuille",
             f"=SUMPRODUCT({_ref(POS, f'$T${L0}:$T${p_liq}')},{_ref(POS, f'$M${L0}:$M${p_liq}')})",
             f"=SUMPRODUCT({_ref(POS, f'$U${L0}:$U${p_liq}')},{_ref(POS, f'$M${L0}:$M${p_liq}')})", "0.0000"),
        11: ("Sensibilité 5 ans du portefeuille",
             f"=SUMPRODUCT({_ref(POS, f'$T${L0}:$T${p_liq}')},{_ref(POS, f'$N${L0}:$N${p_liq}')})",
             f"=SUMPRODUCT({_ref(POS, f'$U${L0}:$U${p_liq}')},{_ref(POS, f'$N${L0}:$N${p_liq}')})", "0.0000"),
        12: ("Sensibilité 10 ans du portefeuille",
             f"=SUMPRODUCT({_ref(POS, f'$T${L0}:$T${p_liq}')},{_ref(POS, f'$O${L0}:$O${p_liq}')})",
             f"=SUMPRODUCT({_ref(POS, f'$U${L0}:$U${p_liq}')},{_ref(POS, f'$O${L0}:$O${p_liq}')})", "0.0000"),
        13: ("Variation moyenne en queue, 2 ans",
             f"=SUMPRODUCT({_ref(HIS, f'$D${L0}:$D${h_fin}')},E{L0}:E{h_fin})",
             f"=SUMPRODUCT({_ref(HIS, f'$D${L0}:$D${h_fin}')},F{L0}:F{h_fin})", TAUX),
        14: ("Variation moyenne en queue, 5 ans",
             f"=SUMPRODUCT({_ref(HIS, f'$E${L0}:$E${h_fin}')},E{L0}:E{h_fin})",
             f"=SUMPRODUCT({_ref(HIS, f'$E${L0}:$E${h_fin}')},F{L0}:F{h_fin})", TAUX),
        15: ("Variation moyenne en queue, 10 ans",
             f"=SUMPRODUCT({_ref(HIS, f'$F${L0}:$F${h_fin}')},E{L0}:E{h_fin})",
             f"=SUMPRODUCT({_ref(HIS, f'$F${L0}:$F${h_fin}')},F{L0}:F{h_fin})", TAUX),
        16: ("VaR (EUR)", f"=LARGE(C{L0}:C{h_fin},I7-ROUNDUP(ROUND(RL_Confiance*I7,9),0)+1)",
             f"=LARGE(D{L0}:D{h_fin},J7-ROUNDUP(ROUND(RL_Confiance*J7,9),0)+1)", EUR),
        17: ("Expected Shortfall (EUR)", f"=SUMPRODUCT(C{L0}:C{h_fin},E{L0}:E{h_fin})",
             f"=SUMPRODUCT(D{L0}:D{h_fin},F{L0}:F{h_fin})", EUR),
        18: ("Scénarios au-delà du seuil", f"=SUMPRODUCT(--(C{L0}:C{h_fin}>I9))", f"=SUMPRODUCT(--(D{L0}:D{h_fin}>J9))", "0"),
        19: ("Scénarios égaux au seuil", f"=SUMPRODUCT(--(C{L0}:C{h_fin}=I9))", f"=SUMPRODUCT(--(D{L0}:D{h_fin}=J9))", "0"),
        20: ("Somme des poids de queue", f"=SUM(E{L0}:E{h_fin})", f"=SUM(F{L0}:F{h_fin})", "0.000000"),
        21: ("Sensibilité spread du portefeuille",
             f"=SUMPRODUCT({_ref(POS, f'$T${L0}:$T${p_liq}')},{_ref(POS, f'$AF${L0}:$AF${p_liq}')})",
             f"=SUMPRODUCT({_ref(POS, f'$U${L0}:$U${p_liq}')},{_ref(POS, f'$AF${L0}:$AF${p_liq}')})", "0.0000"),
        22: ("Variation moyenne en queue, spread",
             f"=SUMPRODUCT({_ref(HIS, f'$G${L0}:$G${h_fin}')},E{L0}:E{h_fin})",
             f"=SUMPRODUCT({_ref(HIS, f'$G${L0}:$G${h_fin}')},F{L0}:F{h_fin})", TAUX),
    }
    for r, (libelle, f_init, f_apres, fmt) in resume.items():
        _cellule(ws, f"H{r}", libelle)
        _cellule(ws, f"I{r}", f_init, fmt)
        _cellule(ws, f"J{r}", f_apres, fmt)
    for i in range(N_SCENARIOS):
        r = L0 + i
        hr = L0 + i
        _cellule(ws, f"B{r}", f"={_ref(HIS, f'B{hr}')}", DATE)
        _cellule(ws, f"C{r}", f"={nav}*($I$10*{_ref(HIS, f'D{hr}')}+$I$11*{_ref(HIS, f'E{hr}')}+$I$12*{_ref(HIS, f'F{hr}')}+$I$21*{_ref(HIS, f'G{hr}')})", EUR)
        _cellule(ws, f"D{r}", f"={nav}*($J$10*{_ref(HIS, f'D{hr}')}+$J$11*{_ref(HIS, f'E{hr}')}+$J$12*{_ref(HIS, f'F{hr}')}+$J$21*{_ref(HIS, f'G{hr}')})", EUR)
        _cellule(ws, f"E{r}", f"=IF(C{r}>$I$9,1,IF(C{r}=$I$9,($I$8-$I$18)/$I$19,0))/$I$8", "0.000")
        _cellule(ws, f"F{r}", f"=IF(D{r}>$J$9,1,IF(D{r}=$J$9,($J$8-$J$18)/$J$19,0))/$J$8", "0.000")
    ws.freeze_panes = f"B{L0}"

    # ------------------------------------------------------------------ Flux
    ws = feuilles[FLX]
    _titre(ws, "Flux du modèle", "Coupons annuels et remboursement à l'échéance, en % du nominal. "
                                 "Valeur actuelle après calibration du spread ; la somme par titre égale le prix dirty.")
    _entete(ws, L0 - 1, 2, ["ISIN", "Date du flux", "Temps (années)", "Montant (% nominal)",
                            "Valeur actuelle (% nominal)"], [14, 12, 12, 14, 16])
    for i, (isin, f) in enumerate(flux):
        r = L0 + i
        _cellule(ws, f"B{r}", isin)
        _cellule(ws, f"C{r}", f.date, DATE)
        _cellule(ws, f"D{r}", f.t, "0.0000", F_SAISIE)
        _cellule(ws, f"E{r}", f.montant, "0.000", F_SAISIE)
        _cellule(ws, f"F{r}", f.valeur_actuelle, "0.000000", F_SAISIE)
    ws.freeze_panes = f"B{L0}"

    # ------------------------------------------------------------------ Backtest
    ws = feuilles[BKT]
    _titre(ws, "Contrôle glissant de la VaR (portefeuille initial)",
           "Chaque VaR n'utilise que les 250 pertes strictement antérieures. P&L modélisé à positions constantes : "
           "ce test vérifie le modèle de taux, pas les pertes réalisées par le fonds.")
    _entete(ws, L0 - 1, 2, ["Date", "Perte (EUR)", "VaR prévue (EUR)", "Exception"], [12, 15, 15, 10])
    _entete(ws, L0 - 1, 7, ["Synthèse", "Valeur"], [34, 14])
    debut = N_SCENARIOS - 250
    for i in range(250):
        r = L0 + i
        cr = L0 + debut + i
        _cellule(ws, f"B{r}", f"={_ref(CAL, f'B{cr}')}", DATE)
        _cellule(ws, f"C{r}", f"={_ref(CAL, f'C{cr}')}", EUR)
        _cellule(ws, f"D{r}", f"=LARGE({_ref(CAL, f'C{cr - 250}:C{cr - 1}')},250-ROUNDUP(ROUND(RL_Confiance*250,9),0)+1)", EUR)
        _cellule(ws, f"E{r}", f"=IF(C{r}>D{r},1,0)", "0")
    b_fin = L0 + 249
    synthese = [
        ("Prévisions", "=COUNT(D7:D256)", "0"),
        ("Exceptions observées", f"=SUM(E{L0}:E{b_fin})", "0"),
        ("Exceptions attendues", "=H7*(1-RL_Confiance)", "0.0"),
        ("Statistique de Kupiec (LR)",
         "=-2*((H7-H8)*LN(RL_Confiance)+H8*LN(1-RL_Confiance)-IF(H8=0,0,(H7-H8)*LN(1-H8/H7)+H8*LN(H8/H7)))", "0.000"),
        ("p-value (khi-deux, 1 degré)", "=CHIDIST(H10,1)", "0.000"),
        ("Conclusion à 5 %", '=IF(H11<0.05,"Calibration rejetée","Calibration non rejetée")', None),
    ]
    for i, (libelle, formule, fmt) in enumerate(synthese):
        _cellule(ws, f"G{7 + i}", libelle)
        _cellule(ws, f"H{7 + i}", formule, fmt)
    ws.conditional_formatting.add(f"E{L0}:E{b_fin}", CellIsRule(operator="equal", formula=["1"], fill=FOND_ALERTE))
    ws.freeze_panes = f"B{L0}"

    # ------------------------------------------------------------------ Inventaire
    ws = feuilles[INV]
    _titre(ws, "Inventaire publié, section taux fixe EUR",
           f"{len(inventaire)} lignes extraites du PDF. Les indicateurs se recalculent avec les seuils de l'onglet Paramètres.")
    _entete(ws, L0 - 1, 2, ["ISIN", "Libellé", "Émetteur", "Échéance", "Coupon (%)", "Nominal (EUR)",
                            "Prix clean", "Coupon couru (EUR)", "Valeur publiée (EUR)", "Date du prix", "Page",
                            "Code entre parenthèses", "Écart de rapprochement (EUR)", "Dans le modèle",
                            "Motif d'exclusion", "Spread calibré (pb)", "Échu", "Valorisation nulle",
                            "Prix ancien", "Spread élevé", "Écart de valorisation", "Suite", "Rang d'alerte"],
            [14, 24, 14, 11, 9, 14, 9, 13, 15, 11, 6, 11, 13, 9, 34, 10, 7, 10, 8, 8, 10, 11, 8])
    pos_isin = _ref(POS, f"$B${L0}:$B${p_fin}")
    pos_spread = _ref(POS, f"$L${L0}:$L${p_fin}")
    for i, l in enumerate(inventaire):
        r = L0 + i
        motifs = motifs_exclusion(l)
        valeurs = [l.isin, l.libelle, l.emetteur, l.echeance, l.coupon, l.nominal, l.prix_clean, l.coupon_couru,
                   l.valeur, l.date_prix, l.page, l.code_parenthese, l.ecart_rapprochement,
                   "NON" if motifs else "OUI", " ; ".join(motifs)]
        fmts = [None, None, None, DATE, "0.000", EUR, "0.0000", EUR2, EUR2, DATE, "0", None, EUR2, None, None]
        for j, (v, fmt) in enumerate(zip(valeurs, fmts)):
            _cellule(ws, f"{get_column_letter(2 + j)}{r}", v, fmt)
        _cellule(ws, f"Q{r}", f'=IFERROR(INDEX({pos_spread},MATCH(B{r},{pos_isin},0)),"")', "0")
        _cellule(ws, f"R{r}", f'=IF(E{r}<=RL_Date,"OUI","")')
        _cellule(ws, f"S{r}", f'=IF(OR(H{r}<=0,J{r}<=0),"OUI","")')
        _cellule(ws, f"T{r}", f'=IF(RL_Date-K{r}>RL_Age_Prix,"OUI","")')
        _cellule(ws, f"U{r}", f'=IF(AND(ISNUMBER(Q{r}),Q{r}>RL_Spread_Revue),"OUI","")')
        _cellule(ws, f"V{r}", f'=IF(ABS(N{r})>{SEUILS["ecart_rapprochement"]},"OUI","")')
        _cellule(ws, f"W{r}", f'=IF(OR(R{r}="OUI",S{r}="OUI",T{r}="OUI",V{r}="OUI"),"À instruire",IF(U{r}="OUI","À examiner",""))')
        _cellule(ws, f"X{r}", f'=IF(W{r}="","",COUNTIF($W${L0}:W{r},"?*"))', "0")
    ws.conditional_formatting.add(f"W{L0}:W{i_fin}", CellIsRule(operator="equal", formula=['"À instruire"'], fill=FOND_ALERTE))
    ws.conditional_formatting.add(f"W{L0}:W{i_fin}", CellIsRule(operator="equal", formula=['"À examiner"'], fill=FOND_EXAMEN))
    ws.freeze_panes = f"D{L0}"

    # ------------------------------------------------------------------ Émetteurs
    ws = feuilles[EMT]
    _titre(ws, "Dispersion par émetteur (règle de type 5/10/40, illustrative)",
           "En % de l'actif net publié. Seule la section taux fixe EUR est connue : ces ratios sont des minimums. "
           "Émetteurs regroupés par libellé abrégé, faute d'identifiant émetteur.")
    emetteurs = sorted({l.emetteur for l in inventaire},
                       key=lambda e: -sum(l.valeur for l in inventaire if l.emetteur == e))
    souverains = {l.emetteur for l in inventaire if l.isin in SOUVERAINS}
    _entete(ws, L0 - 1, 2, ["Émetteur", "Valeur (EUR)", "Poids dans l'actif net", "Souverain", "Statut",
                            "Poids hors souverain"], [22, 16, 14, 10, 26, 12])
    inv_val = _ref(INV, f"$J${L0}:$J${i_fin}")
    inv_emt = _ref(INV, f"$D${L0}:$D${i_fin}")
    for i, e in enumerate(emetteurs):
        r = L0 + i
        _cellule(ws, f"B{r}", e)
        _cellule(ws, f"C{r}", f"=SUMIFS({inv_val},{inv_emt},B{r})", EUR)
        _cellule(ws, f"D{r}", f"=C{r}/RL_Actif_Net", PCT)
        _cellule(ws, f"E{r}", "OUI" if e in souverains else "NON")
        _cellule(ws, f"F{r}", f'=IF(E{r}="OUI","Souverain, régime dérogatoire",IF(D{r}>RL_Seuil_Emetteur_10,"Au-delà du plafond",IF(D{r}>RL_Seuil_Emetteur_5,"Au-delà de 5 %","")))')
        _cellule(ws, f"G{r}", f'=IF(E{r}="NON",D{r},0)', PCT)
    e_fin = L0 + len(emetteurs) - 1
    _entete(ws, L0 - 1, 9, ["Synthèse", "Valeur"], [40, 16])
    synth = [
        ("Plus gros émetteur non souverain", f'=INDEX(B{L0}:B{e_fin},MATCH(MAX(G{L0}:G{e_fin}),G{L0}:G{e_fin},0))', None),
        ("Son poids dans l'actif net", f'=MAX(G{L0}:G{e_fin})', PCT),
        ("Émetteurs non souverains au-delà de 5 %", f'=COUNTIFS(E{L0}:E{e_fin},"NON",D{L0}:D{e_fin},">"&RL_Seuil_Emetteur_5)', "0"),
        ("Cumul de ces émetteurs", f'=SUMIFS(D{L0}:D{e_fin},E{L0}:E{e_fin},"NON",D{L0}:D{e_fin},">"&RL_Seuil_Emetteur_5)', PCT),
        ("Émetteurs au-delà du plafond", f'=COUNTIFS(E{L0}:E{e_fin},"NON",D{L0}:D{e_fin},">"&RL_Seuil_Emetteur_10)', "0"),
        ("Conclusion", '=IF(AND(J11=0,J10<=RL_Seuil_Cumul_40),"Aucun dépassement sur la section connue","À examiner")', None),
    ]
    for i, (libelle, formule, fmt) in enumerate(synth):
        _cellule(ws, f"I{7 + i}", libelle)
        _cellule(ws, f"J{7 + i}", formule, fmt)
    ws.freeze_panes = f"B{L0}"

    # ------------------------------------------------------------------ Contrôles
    ws = feuilles[CTL]
    _titre(ws, "Contrôles et rapprochements",
           "Rapprochements recalculés par formules. Les alertes portent sur l'inventaire publié complet.")
    _entete(ws, 5, 2, ["Contrôle", "Mesure", "Attendu", "Résultat"], [46, 18, 18, 14])
    contributions_ecart = f"={_ref(POS, f'W{p_tot}')}-{_ref(CAL, 'J17')}"
    noeuds_ecart = f"={_ref(COM, 'D29')}+{_ref(COM, 'D30')}+{_ref(COM, 'D31')}+{_ref(COM, 'D32')}-{_ref(CAL, 'J17')}"
    controles = [
        ("Total de l'inventaire = sous-total publié (EUR)", f"=SUM({inv_val})", SOUS_TOTAL_TAUX_FIXE_EUR, EUR2, 0.02),
        ("Somme des poids après vente", f"={_ref(POS, f'U{p_tot}')}", 1, "0.000000", 1e-9),
        ("Contributions par ligne moins ES après vente (EUR)", contributions_ecart, 0, EUR2, 0.01),
        ("Contributions par facteur moins ES après vente (EUR)", noeuds_ecart, 0, EUR2, 0.01),
        ("Somme des poids de queue après vente", f"={_ref(CAL, 'J20')}", 1, "0.000000", 1e-9),
        ("Flux actualisés moins prix dirty (écart maximal)",
         f"=MAX(MAX({_ref(POS, f'AD{L0}:AD{p_fin}')}),-MIN({_ref(POS, f'AD{L0}:AD{p_fin}')}))",
         0, "0.000000000", 1e-6),
    ]
    for i, (libelle, formule, attendu, fmt, tolerance) in enumerate(controles):
        r = 6 + i
        _cellule(ws, f"B{r}", libelle)
        _cellule(ws, f"C{r}", formule, fmt)
        _cellule(ws, f"D{r}", attendu, fmt)
        _cellule(ws, f"E{r}", f'=IF(ABS(C{r}-D{r})<{tolerance},"OK","ÉCART")', police=F_GRAS)
    _nommer(wb, "RL_Ecart_Contributions", CTL, "$C$8")
    c_fin = 6 + len(controles) - 1
    ws.conditional_formatting.add(f"E6:E{c_fin}", CellIsRule(operator="equal", formula=['"ÉCART"'], fill=FOND_ALERTE))
    r = c_fin + 2
    _cellule(ws, f"B{r}", "Dernière vérification VBA indépendante", police=F_GRAS)
    _nommer(wb, "RL_Horodatage", CTL, f"$C${r}")
    _nommer(wb, "RL_Statut_VBA", CTL, f"$E${r}")
    r_alertes = r + 2
    _entete(ws, r_alertes, 2, ["Ligne en alerte", "ISIN", "Suite", "Constat"], None)
    ws.column_dimensions["E"].width = 70
    ws.column_dimensions["F"].width = 50
    inv = lambda col: _ref(INV, f"${col}${L0}:${col}${i_fin}")
    for k in range(1, 26):
        r = r_alertes + k
        m = f"MATCH({k},{inv('X')},0)"
        _cellule(ws, f"B{r}", f'=IFERROR(INDEX({inv("C")},{m}),"")')
        _cellule(ws, f"C{r}", f'=IFERROR(INDEX({inv("B")},{m}),"")')
        _cellule(ws, f"D{r}", f'=IFERROR(INDEX({inv("W")},{m}),"")', police=F_GRAS)
        constat = (f'=IFERROR(TRIM(IF(INDEX({inv("R")},{m})="OUI","Titre échu toujours en inventaire. ","")'
                   f'&IF(INDEX({inv("S")},{m})="OUI","Valorisation nulle. ","")'
                   f'&IF(INDEX({inv("T")},{m})="OUI","Prix daté du "&TEXT(INDEX({inv("K")},{m}),"dd/mm/yyyy")&". ","")'
                   f'&IF(INDEX({inv("U")},{m})="OUI","Spread calibré de "&TEXT(INDEX({inv("Q")},{m}),"0")&" pb. ","")'
                   f'&IF(INDEX({inv("V")},{m})="OUI","Écart de valorisation. ","")),"")')
        _cellule(ws, f"E{r}", constat)
    ws.conditional_formatting.add(f"D{r_alertes + 1}:D{r_alertes + 25}", CellIsRule(operator="equal", formula=['"À instruire"'], fill=FOND_ALERTE))
    ws.conditional_formatting.add(f"D{r_alertes + 1}:D{r_alertes + 25}", CellIsRule(operator="equal", formula=['"À examiner"'], fill=FOND_EXAMEN))
    _cellule(ws, f"B{r_alertes + 27}", "Données indisponibles, hors du périmètre des calculs :", police=F_GRAS)
    for k, (theme, constat, suite) in enumerate(indisponibles):
        r = r_alertes + 28 + k
        _cellule(ws, f"B{r}", theme)
        _cellule(ws, f"C{r}", constat, police=F_NOTE)
        _cellule(ws, f"F{r}", suite, police=F_NOTE)

    # ------------------------------------------------------------------ Comité
    ws = ws_com
    _titre(ws, "RiskLens : synthèse pour le comité des risques",
           f"R-co Conviction Credit Euro, inventaire publié au {DATE_REFERENCE:%d/%m/%Y}. Section obligations à taux "
           f"fixe EUR, {n_pos} lignes modélisées. Risque à un jour sur chocs observés : courbe BCE"
           + (", spread crédit Bundesbank." if spread_ok else " (spread crédit non raccordé)."))
    for col, w in zip("BCDEFGH", [40, 17, 17, 17, 3, 32, 34]):
        ws.column_dimensions[col].width = w
    _entete(ws, 5, 2, ["Indicateur", "Initial", "Après vente", "Écart"])
    lignes_com = [
        ("Valeur du portefeuille modélisé (EUR)", "=RL_NAV", "=RL_NAV", EUR),
        ("Part de l'actif net publié du fonds", "=RL_NAV/RL_Actif_Net", "=RL_NAV/RL_Actif_Net", PCT),
        ("Niveau de confiance", "=RL_Confiance", "=RL_Confiance", PCT),
        ("VaR à 1 jour (EUR)", f"={_ref(CAL, 'I16')}", f"={_ref(CAL, 'J16')}", EUR),
        ("Expected Shortfall à 1 jour (EUR)", f"={_ref(CAL, 'I17')}", f"={_ref(CAL, 'J17')}", EUR),
        ("VaR en % de la valeur", "=C9/RL_NAV", "=D9/RL_NAV", PCT),
        ("Duration du portefeuille", f"={_ref(CAL, 'I10')}+{_ref(CAL, 'I11')}+{_ref(CAL, 'I12')}",
         f"={_ref(CAL, 'J10')}+{_ref(CAL, 'J11')}+{_ref(CAL, 'J12')}", "0.00"),
        ("Stress linéaire (EUR)", f"={_ref(POS, f'Y{p_tot}')}", f"={_ref(POS, f'Z{p_tot}')}", EUR),
        ("Stress par revalorisation des flux (EUR)", f"={_ref(POS, f'AA{p_tot}')}", f"={_ref(POS, f'AB{p_tot}')}", EUR),
        ("Liquidités (EUR)", f"={_ref(POS, f'J{p_liq}')}", f"={_ref(POS, f'S{p_liq}')}", EUR),
    ]
    for i, (libelle, f_init, f_apres, fmt) in enumerate(lignes_com):
        r = 6 + i
        _cellule(ws, f"B{r}", libelle)
        _cellule(ws, f"C{r}", f_init, fmt)
        _cellule(ws, f"D{r}", f_apres, fmt)
        _cellule(ws, f"E{r}", f"=D{r}-C{r}", fmt)
        for col in "BCDE":
            ws[f"{col}{r}"].border = TRAIT
    _nommer(wb, "RL_VaR_Initial", COM, "$C$9")
    _nommer(wb, "RL_VaR_Apres", COM, "$D$9")
    _nommer(wb, "RL_ES_Initial", COM, "$C$10")
    _nommer(wb, "RL_ES_Apres", COM, "$D$10")
    _cellule(ws, "B17", "Stress appliqué", police=F_GRAS)
    _cellule(ws, "C17", '="Taux "&TEXT(RL_Choc_Taux,"+0;-0")&" pb, spreads "&TEXT(RL_Choc_Spread,"+0;-0")&" pb"')

    # Vente simulée
    _entete(ws, 19, 2, ["Vente simulée", "Valeur", "", ""])
    vente = [
        ("Ligne vendue", f"=INDEX({_ref(POS, f'$C${L0}:$C${p_fin}')},MATCH(RL_ISIN_Vente,RL_Liste_ISIN,0))", None),
        ("Montant vendu (EUR)", "=RL_Montant_Vente", EUR),
        ("ES par M EUR détenu sur cette ligne, avant vente",
         f"=INDEX({_ref(POS, f'$V${L0}:$V${p_fin}')},MATCH(RL_ISIN_Vente,RL_Liste_ISIN,0))/"
         f"INDEX({_ref(POS, f'$J${L0}:$J${p_fin}')},MATCH(RL_ISIN_Vente,RL_Liste_ISIN,0))*1000000", EUR),
        ("Réduction d'ES par M EUR vendu", '=IF(RL_Montant_Vente>0,(C10-D10)/RL_Montant_Vente*1000000,0)', EUR),
        ("Référence : vente au prorata, par M EUR vendu", "=C10/RL_NAV*1000000", EUR),
        ("Efficacité relative au prorata", '=IF(C24>0,C23/C24,0)', '0.00"x"'),
    ]
    for i, (libelle, formule, fmt) in enumerate(vente):
        r = 20 + i
        _cellule(ws, f"B{r}", libelle)
        _cellule(ws, f"C{r}", formule, fmt)
        ws[f"C{r}"].alignment = Alignment(horizontal="right")
        for col in "BC":
            ws[f"{col}{r}"].border = TRAIT

    # Contributions par nœud
    _entete(ws, 28, 2, ["Contribution à l'ES par facteur", "Initial (EUR)", "Après vente (EUR)", ""])
    for k, (nom, l_sens, l_choc) in enumerate((("Taux 2 ans", 10, 13), ("Taux 5 ans", 11, 14),
                                                ("Taux 10 ans", 12, 15), ("Spread crédit", 21, 22))):
        r = 29 + k
        _cellule(ws, f"B{r}", nom if (k < 3 or spread_ok) else "Spread crédit (non raccordé)")
        _cellule(ws, f"C{r}", f"=RL_NAV*{_ref(CAL, f'I{l_sens}')}*{_ref(CAL, f'I{l_choc}')}", EUR)
        _cellule(ws, f"D{r}", f"=RL_NAV*{_ref(CAL, f'J{l_sens}')}*{_ref(CAL, f'J{l_choc}')}", EUR)
    graphique = BarChart()
    graphique.type = "bar"
    graphique.title = "Contribution à l'ES par facteur (EUR)"
    graphique.add_data(Reference(ws, min_col=3, max_col=4, min_row=28, max_row=32), titles_from_data=True)
    graphique.set_categories(Reference(ws, min_col=2, min_row=29, max_row=32))
    graphique.height, graphique.width = 6, 12
    ws.add_chart(graphique, "G28")

    # Principaux contributeurs et alertes
    _entete(ws, 5, 7, ["Premiers contributeurs à l'ES", "ES après (EUR)"])
    pos_w = _ref(POS, f"$W${L0}:$W${p_fin}")
    pos_lib = _ref(POS, f"$C${L0}:$C${p_fin}")
    for k in range(1, 11):
        r = 5 + k
        _cellule(ws, f"H{r}", f"=LARGE({pos_w},{k})", EUR)
        _cellule(ws, f"G{r}", f"=INDEX({pos_lib},MATCH(H{r},{pos_w},0))")
        for col in "GH":
            ws[f"{col}{r}"].border = TRAIT
    _entete(ws, 19, 7, ["Contrôles", "Résultat"])
    syntheses = [
        ("Lignes « À instruire »", f'=COUNTIF({_ref(INV, f"$W${L0}:$W${i_fin}")},"À instruire")'),
        ("Lignes « À examiner »", f'=COUNTIF({_ref(INV, f"$W${L0}:$W${i_fin}")},"À examiner")'),
        ("Lignes au-dessus du seuil de revue", f'=COUNTIF({_ref(POS, f"$AC${L0}:$AC${p_fin}")},"OUI")'),
        ("Dispersion émetteurs", f"={_ref(EMT, '$J$12')}"),
        ("Rapprochements en écart", f'=COUNTIF({_ref(CTL, f"$E$6:$E${c_fin}")},"ÉCART")'),
        ("Backtest (Kupiec)", f"={_ref(BKT, '$H$12')}"),
    ]
    for i, (libelle, formule) in enumerate(syntheses):
        r = 20 + i
        _cellule(ws, f"G{r}", libelle)
        _cellule(ws, f"H{r}", formule)
        ws[f"H{r}"].alignment = Alignment(horizontal="right")
        for col in "GH":
            ws[f"{col}{r}"].border = TRAIT
    notes = [
        "Périmètre : section taux fixe EUR de l'inventaire publié ; dérivés, autres sections et passifs exclus. "
        "Ce n'est ni la VaR du fonds ni un contrôle de ses limites.",
        ("Mesure : taux et spread crédit. Le facteur de spread est un indice d'obligations d'entreprises allemandes "
         f"(Bundesbank), bêta par ligne {'proportionnel au spread (DTS)' if methode_spread == 'dts' else 'égal à 1'}."
         if spread_ok else
         "Mesure : risque de taux seulement, faute de série de spread extraite (lancer risklens.extraction). "
         "Le risque de spread n'est mesuré qu'en stress."),
        "Modèle : flux annuels théoriques, courbe AAA BCE interpolée 2/5/10 ans, spread constant ajusté au prix dirty ; "
        "options ignorées. ES à 99 % fondée sur 10 journées équivalentes.",
    ]
    for i, texte in enumerate(notes):
        _cellule(ws, f"B{42 + i}", texte, police=F_NOTE)
    ws.print_area = "B2:H45"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _nommer(wb, "RL_Zone_Comite", COM, "$B$2:$H$45")

    # ------------------------------------------------------------------ Sources
    ws = feuilles[SRC]
    _titre(ws, "Sources, méthode et traçabilité")
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110
    manifeste = charger_manifeste()
    src = manifeste["sources"]
    lignes_src = [
        ("Inventaire", src["inventaire_pdf"]["description"] + " (Rothschild & Co Asset Management)"),
        ("Adresse du PDF", src["inventaire_pdf"]["url"]),
        ("SHA-256 du PDF", src["inventaire_pdf"]["sha256"]),
        ("Courbes de taux", "BCE, courbe zéro-coupon AAA zone euro : YC.B.U2.EUR.4F.G_N_A.SV_C_YM.SR_2Y, SR_5Y, SR_10Y"),
        ("Spread crédit", "Bundesbank : BBSIS.D.I.UMR.RD.EUR.X2000 (obligations d'entreprises, Nicht-MFIs) moins "
                          "BBSIS.D.I.UMR.RD.EUR.S1311.B.A604 (titres fédéraux cotés), séries quotidiennes"
                          + ("" if spread_ok else " ; non extrait dans cette version")),
        ("Méthode de spread", {"dts": "Bêta proportionnel au spread calibré (Duration Times Spread)",
                               "uniforme": "Bêta égal à 1 pour toutes les lignes", None: "Sans objet"}[methode_spread]),
        ("Manifeste généré le", manifeste["genere_le"][:10]),
    ] + [(f"SHA-256 {nom}", valeur) for nom, valeur in empreintes.items()] + [
        ("Univers modélisé", "Toutes les lignes de la section taux fixe EUR valorisées au 30/06/2025, à prix positif, "
                             "d'échéance comprise entre la date de référence et 12 ans. Aucun filtre sur le code entre parenthèses."),
        ("Lignes exclues", f"{len(inventaire) - n_pos} ligne(s), listées avec leur motif dans l'onglet Inventaire."),
        ("Calibration", "Spread continu constant ajusté par dichotomie au prix dirty publié (valeur / nominal)."),
        ("Sensibilités", "Dérivées du prix par rapport aux nœuds 2, 5 et 10 ans ; taux constant au-delà de 10 ans."),
        ("VaR et ES", "VaR : perte au rang arrondi supérieur de confiance x N. ES : moyenne pondérée des N x (1 - confiance) pires pertes."),
        ("Contributions", "Allocation d'Euler sur les mêmes journées de queue que l'ES ; la somme égale l'ES."),
        ("Généré le", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")),
        ("Avertissement", "Projet personnel, sans lien avec Rothschild & Co. Aucune recommandation d'investissement."),
    ]
    for i, (cle, valeur) in enumerate(lignes_src):
        _cellule(ws, f"B{5 + i}", cle, police=F_GRAS)
        _cellule(ws, f"C{5 + i}", valeur)

    for nom, ws in feuilles.items():
        if nom == COM:
            continue
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(chemin)
